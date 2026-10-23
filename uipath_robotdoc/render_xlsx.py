from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

HEADER_FILL = PatternFill("solid", fgColor="1F4E79")


def _sheet(wb, title, header, rows):
    ws = wb.create_sheet(title)
    ws.append(header)
    for c in ws[1]:
        c.font, c.fill = Font(bold=True, color="FFFFFF"), HEADER_FILL
    for r in rows:
        ws.append(list(r))
    for row in ws.iter_rows(min_row=2):
        for c in row:
            if isinstance(c.value, str) and c.value[:1] in "=+-@":
                c.data_type = "s"  # keep text such as "=foo" from becoming a formula
            c.alignment = Alignment(vertical="top", wrap_text=True)
    for i, col in enumerate(ws.columns, 1):
        width = max((len(str(c.value)) for c in col if c.value is not None), default=8)
        ws.column_dimensions[get_column_letter(i)].width = min(max(width + 2, 10), 70)
    ws.freeze_panes = "A2"
    if ws.max_row > 1:
        ws.auto_filter.ref = ws.dimensions
    return ws


def render(model: dict, out: Path, graph: Path | None = None, mix: Path | None = None) -> Path:
    p, s, wfs = model["package"], model["stats"], model["workflows"]
    wb = Workbook()
    wb.remove(wb.active)
    _sheet(wb, "Summary", ["Item", "Value"], [
        ("Package", p["name"]), ("Version", p["version"]), ("Description", p["description"]),
        ("Authors", p["authors"]), ("Entry point", p["main"]), ("Expression language", p["expression_language"]),
        ("Target framework", p["target_framework"]), ("Workflows", s["workflows"]), ("Activities", s["activities"]),
        ("Arguments", s["arguments"]), ("Variables", s["variables"]), ("Invocations", s["invocations"]),
        ("Dependencies", s["dependencies"]),
    ])
    _sheet(wb, "Workflows", ["Path", "Main", "Root", "Activities", "Arguments", "Variables", "Invokes", "Invoked by", "Notes"],
           [(w["path"], "yes" if w["is_main"] else "", w["root_type"], len(w["activities"]), len(w["arguments"]),
             len(w["variables"]), ", ".join(i["target_path"] for i in w["invokes"]), ", ".join(w["invoked_by"]),
             w["parse_error"] or w["annotation"]) for w in wfs])
    _sheet(wb, "Arguments", ["Workflow", "Name", "Direction", "Type"],
           [(w["path"], a["name"], a["direction"], a["type"]) for w in wfs for a in w["arguments"]])
    _sheet(wb, "Variables", ["Workflow", "Name", "Type", "Default", "Scope"],
           [(w["path"], v["name"], v["type"], v["default_value"], v["scope"]) for w in wfs for v in w["variables"]])
    _sheet(wb, "Activities", ["Workflow", "#", "Depth", "Type", "Display name", "Detail", "Annotation"],
           [(w["path"], a["seq"], a["depth"], a["type"], a["display_name"], a["detail"], a["annotation"])
            for w in wfs for a in w["activities"]])
    _sheet(wb, "Invocations", ["Caller", "Target", "Resolved to"],
           [(w["path"], i["target_path"], i["callee"] or "NOT FOUND") for w in wfs for i in w["invokes"]])
    _sheet(wb, "Dependencies", ["Package", "Version"], [(d["name"], d["version"]) for d in model["dependencies"]])
    wb.save(out)
    return out
