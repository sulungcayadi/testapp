from __future__ import annotations

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Inches, Pt

NAVY = RGBColor(0x1F, 0x4E, 0x79)
MAX_ROWS = 12


def _title(slide, text, sub=None):
    tb = slide.shapes.add_textbox(Inches(0.5), Inches(0.3), Inches(12.3), Inches(0.8))
    r = tb.text_frame.paragraphs[0].add_run()
    r.text, r.font.size, r.font.bold, r.font.color.rgb = text, Pt(28), True, NAVY
    if sub:
        p = tb.text_frame.add_paragraph()
        p.text, p.font.size = sub, Pt(13)


def _table(slide, header, rows, top=1.5, widths=None):
    rows = rows[:MAX_ROWS]
    shape = slide.shapes.add_table(len(rows) + 1, len(header), Inches(0.5), Inches(top), Inches(12.3), Inches(0.35 * (len(rows) + 1)))
    t = shape.table
    for i, h in enumerate(header):
        t.cell(0, i).text = h
    for ri, r in enumerate(rows, 1):
        for ci, v in enumerate(r):
            t.cell(ri, ci).text = "" if v is None else str(v)[:90]
    for row in t.rows:
        for c in row.cells:
            for para in c.text_frame.paragraphs:
                para.font.size = Pt(11)
    return t


def _picture(slide, path: Path, top=1.3, max_w=12.3, max_h=5.8):
    from PIL import Image
    w, h = Image.open(path).size
    scale = min(max_w / (w / 100), max_h / (h / 100))
    width, height = w / 100 * scale, h / 100 * scale
    slide.shapes.add_picture(str(path), Inches((13.333 - width) / 2), Inches(top), Inches(width), Inches(height))


def render(model: dict, out: Path, graph: Path | None, mix: Path | None) -> Path:
    p, s = model["package"], model["stats"]
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    blank = prs.slide_layouts[6]

    sl = prs.slides.add_slide(blank)
    _title(sl, p["name"], f"Robot documentation  |  version {p['version'] or '-'}")
    tb = sl.shapes.add_textbox(Inches(0.5), Inches(2.2), Inches(12), Inches(2))
    tb.text_frame.word_wrap = True
    tb.text_frame.text = p["description"] or "No description provided."
    tb.text_frame.paragraphs[0].font.size = Pt(18)

    sl = prs.slides.add_slide(blank)
    _title(sl, "Overview")
    _table(sl, ["Item", "Value"], [
        ("Entry point", p["main"] or "-"), ("Expression language", p["expression_language"] or "-"),
        ("Target framework", p["target_framework"] or "-"), ("Workflows", s["workflows"]), ("Activities", s["activities"]),
        ("Arguments", s["arguments"]), ("Variables", s["variables"]), ("Dependencies", s["dependencies"])])

    for img, name in ((graph, "Workflow call graph"), (mix, "Activity usage")):
        if img and img.exists():
            sl = prs.slides.add_slide(blank)
            _title(sl, name)
            _picture(sl, img)

    if model["dependencies"]:
        sl = prs.slides.add_slide(blank)
        _title(sl, "Dependencies")
        _table(sl, ["Package", "Version"], [(d["name"], d["version"]) for d in model["dependencies"]])

    for w in model["workflows"]:
        sl = prs.slides.add_slide(blank)
        _title(sl, w["name"] + (" (main)" if w["is_main"] else ""),
               f"{w['path']}  |  {len(w['activities'])} activities  |  root {w['root_type'] or '-'}")
        top = 1.6
        if w["arguments"]:
            _table(sl, ["Argument", "Direction", "Type"], [(a["name"], a["direction"], a["type"]) for a in w["arguments"]], top=top)
            top += 0.35 * (min(len(w["arguments"]), MAX_ROWS) + 1) + 0.3
        if w["invokes"] and top < 6:
            tb = sl.shapes.add_textbox(Inches(0.5), Inches(top), Inches(12.3), Inches(0.6))
            tb.text_frame.word_wrap = True
            tb.text_frame.text = "Invokes: " + ", ".join(i["target_path"] for i in w["invokes"])
            tb.text_frame.paragraphs[0].font.size = Pt(13)
        elif not w["arguments"] and not w["invokes"]:
            tb = sl.shapes.add_textbox(Inches(0.5), Inches(top), Inches(12.3), Inches(0.6))
            tb.text_frame.text = w["annotation"] or "No arguments or invoked workflows."
            tb.text_frame.paragraphs[0].font.size = Pt(14)
    prs.save(out)
    return out
