from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.shared import Inches, Pt

MAX_OUTLINE = 400


def _table(doc, header, rows, widths=None):
    t = doc.add_table(rows=1, cols=len(header))
    t.style = "Light Grid Accent 1"
    for i, h in enumerate(header):
        t.rows[0].cells[i].text = h
    for r in rows:
        cells = t.add_row().cells
        for i, v in enumerate(r):
            cells[i].text = "" if v is None else str(v)
    for row in t.rows:
        for c in row.cells:
            for para in c.paragraphs:
                for run in para.runs:
                    run.font.size = Pt(9)
    doc.add_paragraph()
    return t


def render(model: dict, out: Path, graph: Path | None, mix: Path | None) -> Path:
    p, s = model["package"], model["stats"]
    doc = Document()
    doc.add_heading(f"{p['name']} - Robot Documentation", 0)
    doc.add_paragraph(f"Version {p['version'] or '-'}  |  {p['description'] or 'No description provided.'}")
    doc.add_heading("Overview", 1)
    _table(doc, ["Item", "Value"], [
        ("Entry point", p["main"] or "-"), ("Expression language", p["expression_language"] or "-"),
        ("Target framework", p["target_framework"] or "-"), ("Authors", p["authors"] or "-"),
        ("Workflows", s["workflows"]), ("Activities", s["activities"]), ("Arguments", s["arguments"]),
        ("Variables", s["variables"]), ("Dependencies", s["dependencies"]),
    ])
    notes = ([f"{u} is not invoked by any other workflow." for u in s["unreferenced"]]
             + [f"{a} invokes {b}, which is not in the package." for a, b in s["unresolved_invokes"]]
             + [f"{a} could not be parsed: {b}" for a, b in s["parse_errors"]])
    if notes:
        doc.add_heading("Review notes", 2)
        for n in notes:
            doc.add_paragraph(n, style="List Bullet")
    for img in (graph, mix):
        if img and img.exists():
            doc.add_picture(str(img), width=Inches(6.2))
    doc.add_heading("Dependencies", 1)
    if model["dependencies"]:
        _table(doc, ["Package", "Version"], [(d["name"], d["version"]) for d in model["dependencies"]])
    else:
        doc.add_paragraph("None declared.")
    for w in model["workflows"]:
        doc.add_heading(w["path"] + (" (main)" if w["is_main"] else ""), 1)
        if w["parse_error"]:
            doc.add_paragraph(f"Parse error: {w['parse_error']}")
        if w["annotation"]:
            doc.add_paragraph(w["annotation"])
        doc.add_paragraph(f"Root: {w['root_type'] or '-'}  |  {len(w['activities'])} activities"
                          + (f"  |  invoked by {', '.join(w['invoked_by'])}" if w["invoked_by"] else ""))
        if w["arguments"]:
            doc.add_heading("Arguments", 3)
            _table(doc, ["Name", "Direction", "Type"], [(a["name"], a["direction"], a["type"]) for a in w["arguments"]])
        if w["variables"]:
            doc.add_heading("Variables", 3)
            _table(doc, ["Name", "Type", "Default", "Scope"],
                   [(v["name"], v["type"], v["default_value"], v["scope"]) for v in w["variables"]])
        if w["invokes"]:
            doc.add_heading("Invokes", 3)
            for i in w["invokes"]:
                doc.add_paragraph(i["target_path"] + ("" if i["callee"] else " (not found)"), style="List Bullet")
        if w["activities"]:
            doc.add_heading("Activity outline", 3)
            shown = w["activities"][:MAX_OUTLINE]
            _table(doc, ["#", "Activity", "Type", "Detail"],
                   [(a["seq"], "  " * a["depth"] + (a["display_name"] or ""), a["type"], a["detail"]) for a in shown])
            if len(w["activities"]) > MAX_OUTLINE:
                doc.add_paragraph(f"... {len(w['activities']) - MAX_OUTLINE} more activities (see the xlsx/html output).")
    doc.save(out)
    return out
