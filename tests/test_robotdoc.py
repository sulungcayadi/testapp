import sqlite3
import zipfile

import pytest
from docx import Document
from openpyxl import load_workbook
from pptx import Presentation

from uipath_robotdoc import cli, parser
from uipath_robotdoc.sample import write_sample_nupkg


@pytest.fixture()
def nupkg(tmp_path):
    return write_sample_nupkg(tmp_path / "InvoiceBot.nupkg")


def test_parse_nupkg(nupkg):
    pkg = parser.read_input(nupkg)
    assert (pkg.name, pkg.version, pkg.main) == ("InvoiceBot", "1.0.3", "Main.xaml")
    assert [n for n, _ in pkg.dependencies] == ["UiPath.System.Activities", "UiPath.UIAutomation.Activities"]
    main = next(w for w in pkg.workflows if w.path == "Main.xaml")
    assert [(a.name, a.direction, a.type) for a in main.arguments] == [
        ("in_Config", "In", "Dictionary(String, Object)"), ("out_Status", "Out", "String")]
    assert [(v.name, v.type, v.default) for v in main.variables] == [("retryCount", "Int32", "3"), ("status", "String", "")]
    assert main.root_type == "Sequence" and main.annotation.startswith("Entry point")
    assert main.invokes == ["Framework\\Init.xaml", "Process.xaml", "Missing.xaml"]


def test_single_xaml_and_folder(nupkg, tmp_path):
    with zipfile.ZipFile(nupkg) as zf:
        zf.extractall(tmp_path / "proj")
    pkg = parser.read_input(tmp_path / "proj" / "content")
    assert len(pkg.workflows) == 4
    one = parser.read_input(tmp_path / "proj" / "content" / "Process.xaml")
    assert [w.name for w in one.workflows] == ["Process"]


def test_rejects_entities_and_bad_input(tmp_path):
    wf = parser.parse_xaml(b'<!DOCTYPE a [<!ENTITY x "y">]><Activity/>', "Evil.xaml")
    assert "DOCTYPE" in wf.parse_error and not wf.activities
    assert "XML parse error" in parser.parse_xaml(b"<Activity>", "Broken.xaml").parse_error
    junk = tmp_path / "x.txt"
    junk.write_text("hi")
    with pytest.raises(parser.PackageError):
        parser.read_input(junk)


def test_end_to_end_all_formats(nupkg, tmp_path):
    out = tmp_path / "out"
    assert cli.main([str(nupkg), "-o", str(out)]) == 0
    base = str(out / "InvoiceBot_1.0.3")
    for ext in (".html", ".docx", ".pptx", ".xlsx", ".pdf", "_callgraph.png", "_activities.png"):
        assert (out / f"InvoiceBot_1.0.3{ext}").stat().st_size > 1000, ext
    html = open(base + ".html", encoding="utf-8").read()
    assert "Missing.xaml, which is not in the package" in html and "Framework/Unused.xaml is not invoked" in html
    assert (out / "InvoiceBot_1.0.3.pdf").read_bytes().startswith(b"%PDF")
    wb = load_workbook(base + ".xlsx")
    assert {"Summary", "Workflows", "Arguments", "Variables", "Activities", "Invocations", "Dependencies"} <= set(wb.sheetnames)
    assert any(r[2] == "NOT FOUND" for r in wb["Invocations"].iter_rows(min_row=2, values_only=True))
    assert "Main.xaml" in "\n".join(p.text for p in Document(str(base + ".docx")).paragraphs)
    assert len(Presentation(str(base + ".pptx")).slides) >= 6
    rows = sqlite3.connect(out / "robot_docs.db").execute("SELECT COUNT(*) FROM workflows").fetchone()
    assert rows[0] == 4


def test_format_selection_and_reingest(nupkg, tmp_path):
    out = tmp_path / "out"
    assert cli.main([str(nupkg), "-o", str(out), "-f", "xlsx"]) == 0
    assert cli.main([str(nupkg), "-o", str(out), "-f", "xlsx"]) == 0  # second run replaces, not duplicates
    assert sorted(p.name for p in out.iterdir()) == ["InvoiceBot_1.0.3.xlsx", "robot_docs.db"]
    assert sqlite3.connect(out / "robot_docs.db").execute("SELECT COUNT(*) FROM packages").fetchone()[0] == 1


def test_formula_like_text_stays_text(tmp_path):
    from uipath_robotdoc import render_xlsx
    model = {"package": dict(name="p", version="", description="=1+1", authors="", main="", expression_language="", target_framework=""),
             "stats": dict(workflows=0, activities=0, arguments=0, variables=0, invocations=0, dependencies=0),
             "workflows": [], "dependencies": []}
    render_xlsx.render(model, tmp_path / "f.xlsx")
    ws = load_workbook(tmp_path / "f.xlsx")["Summary"]
    assert ws["B4"].value == "=1+1" and ws["B4"].data_type == "s"
